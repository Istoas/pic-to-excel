import streamlit as st
import google.generativeai as genai
from PIL import Image
import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

# --- CONFIGURATION SÉCURISÉE ---
# On récupère la clé depuis le coffre-fort (Secrets)
# Ça marche automatiquement sur le Cloud ET sur ton PC (si tu as fait l'étape 3)
try:
    api_key = st.secrets["GEMINI_API_KEY"]
    genai.configure(api_key=api_key)
except FileNotFoundError:
    st.error("Erreur : La clé API n'est pas configurée dans les Secrets.")
    st.stop()

# --- DESIGN PAGE (MODE WIDE OBLIGATOIRE POUR FAIRE PRO) ---
st.set_page_config(page_title="Pic-to-Excel", page_icon="📊", layout="wide")

# --- CSS CACHÉ POUR NETTOYER L'INTERFACE ---
st.markdown("""
<style>
    /* Cacher le menu hamburger et le footer "Made with Streamlit" */
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    
    /* Espacement plus aéré */
    .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
    }
</style>
""", unsafe_allow_html=True)

# --- FONCTION EXCEL ---
def create_styled_excel(df):
    buffer = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    
    # Nettoyage des noms de colonnes
    new_columns = []
    for i, col in enumerate(df.columns):
        col_str = str(col).strip()
        if "Unnamed" in col_str or col_str == "":
            new_columns.append(f"Col_{i+1}")
        else:
            new_columns.append(col_str)
    df.columns = new_columns

    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)

    max_row = ws.max_row
    max_col = ws.max_column
    if max_row >= 2 and max_col >= 1:
        last_col_letter = get_column_letter(max_col)
        ref = f"A1:{last_col_letter}{max_row}"
        tab = Table(displayName="TabDonnees", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        tab.tableStyleInfo = style
        ws.add_table(tab)

        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 20 # Largeur par défaut propre
            
    wb.save(buffer)
    return buffer.getvalue()

# --- SIDEBAR (PARAMÈTRES) ---
with st.sidebar:
    st.header("⚙️ Paramètres IA")
    choix_modele = st.radio(
        "Modèle :",
        ["Flash ⚡ (Rapide)", "Pro 🎯 (Précision Max)"],
        index=1, # Pro par défaut pour éviter les erreurs de colonnes
        help="Si l'IA décale les colonnes, utilise le mode Pro."
    )
    
    MODEL_NAME = "gemini-2.5-flash" if "Flash" in choix_modele else "gemini-2.5-pro"
    
    st.divider()
    st.info(f"Moteur actif : {MODEL_NAME}")
    try:
        model = genai.GenerativeModel(MODEL_NAME)
    except:
        st.error("Erreur API")

# --- HEADER ---
st.title("📊 Pic-to-Excel")
st.markdown("**L'outil de conversion intelligent.** Transforme tes images en Excel sans re-saisir une seule case.")
st.divider()

# --- INTERFACE PRINCIPALE (LAYOUT 2 COLONNES) ---
# Gauche : Input (30%) | Droite : Output (70%)
col_gauche, col_droite = st.columns([1, 2], gap="large")

if 'df_result' not in st.session_state:
    st.session_state.df_result = None

with col_gauche:
    st.subheader("1. Importation")
    uploaded_file = st.file_uploader("Dépose ton document ici", type=["jpg", "png", "jpeg", "webp"])
    
    if uploaded_file:
        image = Image.open(uploaded_file)
        # On affiche l'image en petit pour ne pas prendre toute la place
        st.image(image, caption="Aperçu", use_column_width=True)
        
        st.write("") # Espace
        if st.button("⚡ Lancer l'extraction", type="primary", use_container_width=True):
            with st.spinner("Analyse de la structure géométrique..."):
                try:
                    # --- LE PROMPT "GÉOMÈTRE" ---
                    # C'est ici qu'on règle le problème des colonnes
                    prompt = """
                    Tu es un expert en OCR et structure de données.
                    TACHE : Analyse cette image et convertis le tableau en CSV (séparateur point-virgule).
                    
                    RÈGLES D'ALIGNEMENT STRICTES :
                    1. Regarde l'alignement VERTICAL des colonnes.
                    2. Si une case est vide visuellement (pas de texte), tu DOIS insérer un champ vide (;;) dans le CSV. Ne décale pas les données vers la gauche.
                    3. Si une ligne contient des totaux ou des sous-totaux, garde-les.
                    4. Sois précis sur les chiffres (attention aux virgules/points).
                    5. Pas de texte avant ou après le CSV.
                    """
                    
                    response = model.generate_content([prompt, image])
                    raw_csv = response.text.replace("```csv", "").replace("```", "").strip()
                    
                    # Lecture plus souple
                    st.session_state.df_result = pd.read_csv(
                        io.StringIO(raw_csv), 
                        sep=";", 
                        engine="python",
                        on_bad_lines='skip' # On saute les lignes cassées plutôt que de crasher
                    )
                    st.toast("Tableau extrait !", icon="✅")
                    
                except Exception as e:
                    st.error(f"Erreur : {e}")

with col_droite:
    if st.session_state.df_result is not None:
        st.subheader("2. Résultat Éditable")
        
        # L'éditeur prend toute la place disponible
        edited_df = st.data_editor(
            st.session_state.df_result,
            num_rows="dynamic",
            use_container_width=True,
            height=600 # Grand tableau confortable
        )
        
        st.subheader("3. Export")
        excel_data = create_styled_excel(edited_df)
        
        c1, c2, c3 = st.columns([1, 1, 2])
        with c1:
            st.download_button("📥 Excel (.xlsx)", excel_data, "Tableau.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
        with c2:
            st.download_button("📄 CSV Brut", edited_df.to_csv(index=False, sep=";").encode('utf-8'), "data.csv", "text/csv")
    
    elif not uploaded_file:
        # Message d'accueil quand rien n'est chargé (pour remplir le vide)
        st.info("👈 Commence par charger une image dans la colonne de gauche.")
        st.markdown("""
        **Conseils pour une meilleure précision :**
        * Prends la photo bien à plat (pas de travers).
        * Assure-toi que l'éclairage est bon.
        * Le mode **Pro** est plus lent mais gère mieux les colonnes vides.
        """)